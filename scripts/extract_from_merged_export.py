"""Extract per-category raw SellerSprite exports + employee/category
ownership data from a merged multi-category export (a "合并总表" --
many per-category SellerSprite exports concatenated into one file).

This does NOT normalize, classify, or resolve anything -- it only
slices real rows into the same shape a single-category SellerSprite
export already has (same header row, same column meanings), so the
existing scripts/ingest.py path picks them up completely unchanged.
Streams the source file with openpyxl's read_only mode (never loads it
fully into memory) since these merged exports can be very large.

Usage:
    python scripts/extract_from_merged_export.py <path_to_merged_xlsx>

Category mapping (raw Chinese 二级类目 label -> our category_id) is
the one piece of real judgment this script makes -- kept as an
explicit, visible constant below, not inferred, so every category this
project has ever onboarded is traceable to a real decision.
"""

from __future__ import annotations

import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.workbook import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]

# raw 二级类目 (secondary category) label -> our internal category_id.
# Only categories actually onboarded (real taxonomy authored, real
# config/categories.yaml entry added) belong here -- see
# docs/data_dictionary.md and config/categories.yaml for each one's
# real description.
CATEGORY_MAP = {
    "牙科模型": "dental_models",
    "种植体": "implants",
}

RAW_HEADER_COLUMN = "二级类目"
PRODUCT_MANAGER_COLUMN = "产品经理"
SHEET_NAME = "Sheet1"


def extract(source_path: Path) -> None:
    wb = openpyxl.load_workbook(str(source_path), read_only=True, data_only=True)
    ws = wb[SHEET_NAME]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    col_idx = {name: i for i, name in enumerate(header)}

    category_col = col_idx[RAW_HEADER_COLUMN]
    pm_col = col_idx[PRODUCT_MANAGER_COLUMN]

    category_rows: dict[str, list[tuple]] = defaultdict(list)
    ownership_counts: dict[tuple[str, str], int] = Counter()
    total_rows = 0

    for row in rows:
        total_rows += 1
        raw_category = row[category_col]
        product_manager = row[pm_col]

        if raw_category:
            if product_manager:
                ownership_counts[(raw_category, product_manager)] += 1
            if raw_category in CATEGORY_MAP:
                category_rows[raw_category].append(row)

    wb.close()

    print(f"scanned {total_rows} rows")

    for raw_category, category_id in CATEGORY_MAP.items():
        rows_for_category = category_rows.get(raw_category, [])
        if not rows_for_category:
            print(f"WARNING: no rows found for '{raw_category}' ({category_id}) -- skipped")
            continue
        out_dir = PROJECT_ROOT / "data" / "raw" / category_id
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / f"{category_id}_sellersprite.xlsx"
        _write_rows(header, rows_for_category, out_path)
        print(f"{category_id}: {len(rows_for_category)} rows -> {out_path}")

    _write_ownership_csv(ownership_counts)


def _write_rows(header: list, rows: list[tuple], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(header)
    for row in rows:
        ws.append(list(row))
    wb.save(str(out_path))


def _write_ownership_csv(ownership_counts: dict[tuple[str, str], int]) -> None:
    """Real evidence artifact for Milestone 12 (Personnel Intelligence):
    every (raw category label, product manager name, listing count)
    combination found in the source file. Kept immutable under
    data/raw/ like any other raw export, and deliberately keeps the RAW
    Chinese category label rather than our internal category_id --
    most of these 154 raw categories have no taxonomy/onboarding in
    this project yet, and mapping them to a category_id that doesn't
    really exist here would be exactly the kind of fabrication
    PRINCIPLES.md forbids. src/dmie/database/repository.py's loader
    resolves category_id only for the subset that IS onboarded
    (CATEGORY_MAP above)."""
    out_path = PROJECT_ROOT / "data" / "raw" / "employee_category_ownership.csv"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8-sig", newline="") as f:
        f.write("raw_category_label,product_manager,listing_count,matched_category_id\n")
        for (raw_category, product_manager), count in sorted(ownership_counts.items(), key=lambda kv: -kv[1]):
            matched = CATEGORY_MAP.get(raw_category, "")
            raw_category_escaped = raw_category.replace('"', '""')
            product_manager_escaped = product_manager.replace('"', '""')
            f.write(f'"{raw_category_escaped}","{product_manager_escaped}",{count},{matched}\n')
    print(f"employee/category ownership: {len(ownership_counts)} (category, manager) pairs -> {out_path}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python scripts/extract_from_merged_export.py <path_to_merged_xlsx>")
        sys.exit(2)
    extract(Path(sys.argv[1]))
